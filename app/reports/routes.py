import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from flask import Blueprint, Response, flash, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from ..extensions import db
from ..models import Category, Expense, User
from ..utils import VALID_SOURCES, parse_iso_date, period_range, source_label, visible_expense_query


reports_bp = Blueprint("reports", __name__)


def _report_data(include_all_items=False):
    mode = request.args.get("mode", "month")
    if mode not in {"day", "week", "month", "year", "custom"}:
        mode = "month"
    start_arg = parse_iso_date(request.args.get("start", ""))
    end_arg = parse_iso_date(request.args.get("end", ""))
    if mode == "custom" and request.args.get("start") and not start_arg:
        flash("Ngày bắt đầu không hợp lệ.", "warning")
    if mode == "custom" and request.args.get("end") and not end_arg:
        flash("Ngày kết thúc không hợp lệ.", "warning")
    start, end = period_range(mode, start_arg, end_arg)
    if start > end:
        start, end = end, start
    query = visible_expense_query().filter(Expense.date >= start, Expense.date <= end)
    member_value = request.args.get("member", "")
    if current_user.is_admin and member_value.isdigit():
        query = query.filter(Expense.idUser == int(member_value))
    source = request.args.get("source", "")
    if source in VALID_SOURCES:
        query = query.filter(Expense.source == source)
    ordered_query = query.order_by(Expense.date.asc(), Expense.idExpense.asc())
    total_count = ordered_query.count()
    pagination = None if include_all_items else ordered_query.paginate(page=max(request.args.get("page", 1, type=int), 1), per_page=25, error_out=False)
    items = ordered_query.all() if include_all_items else pagination.items

    totals_by_type = {
        category_type: Decimal(str(amount or 0))
        for category_type, amount in query.with_entities(Category.type, db.func.sum(Expense.amount)).group_by(Category.type).all()
    }
    income_total = totals_by_type.get("income", Decimal("0"))
    expense_total = totals_by_type.get("expense", Decimal("0"))
    category_rows = [
        {"name": name, "type": category_type, "amount": Decimal(str(amount or 0))}
        for name, category_type, amount in query.with_entities(
            Category.nameCategory, Category.type, db.func.sum(Expense.amount)
        ).group_by(Category.nameCategory, Category.type).order_by(db.func.sum(Expense.amount).desc()).all()
    ]
    daily_rows = query.filter(Category.type == "expense").with_entities(
        Expense.date, db.func.sum(Expense.amount)
    ).group_by(Expense.date).order_by(Expense.date.asc()).all()
    chart_labels = [day.isoformat() for day, _amount in daily_rows]
    chart_values = [float(amount or 0) for _day, amount in daily_rows]
    members = User.query.order_by(User.hoten.asc()).all() if current_user.is_admin else []
    return {
        "mode": mode,
        "start": start,
        "end": end,
        "items": items,
        "pagination": pagination,
        "total_count": total_count,
        "income_total": income_total,
        "expense_total": expense_total,
        "balance": income_total - expense_total,
        "category_rows": category_rows,
        "chart_labels": chart_labels,
        "chart_values": chart_values,
        "members": members,
        "selected_member": member_value,
        "source": source,
    }


@reports_bp.route("/")
@login_required
def index():
    return render_template("reports/index.html", report=_report_data())


@reports_bp.route("/export.xlsx")
@login_required
def export_xlsx():
    report = _report_data(include_all_items=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Tong quan"
    summary.append(["BÁO CÁO QUẢN LÝ CHI TIÊU TẬP THỂ / TỔ CHỨC"])
    summary.append(["Từ ngày", report["start"].strftime("%d/%m/%Y")])
    summary.append(["Đến ngày", report["end"].strftime("%d/%m/%Y")])
    summary.append(["Tổng thu (VNĐ)", float(report["income_total"])])
    summary.append(["Tổng chi (VNĐ)", float(report["expense_total"])])
    summary.append(["Chênh lệch thu - chi (VNĐ)", float(report["balance"])])
    summary["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="173F5F")
    summary.merge_cells("A1:B1")
    for row in range(4, 7):
        summary.cell(row=row, column=2).number_format = '#,##0.00'
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 24

    sheet = workbook.create_sheet("Giao dich")
    headers = ["Mã giao dịch", "Ngày", "Thành viên", "Loại", "Danh mục", "Nguồn tiền", "Số tiền (VNĐ)", "Ghi chú"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F5F")
        cell.alignment = Alignment(horizontal="center")
    for item in report["items"]:
        sheet.append([
            item.idExpense,
            item.date.strftime("%d/%m/%Y"),
            item.user.hoten,
            "Thu nhập" if item.category.type == "income" else "Chi tiêu",
            item.category.nameCategory,
            source_label(item.source),
            float(item.amount),
            item.purpose,
        ])
        sheet.cell(row=sheet.max_row, column=7).number_format = '#,##0.00'
    widths = [16, 14, 24, 15, 22, 16, 20, 42]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"bao_cao_chi_tieu_{report['start'].isoformat()}_{report['end'].isoformat()}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@reports_bp.route("/export.csv")
@login_required
def export_csv():
    report = _report_data(include_all_items=True)
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Mã giao dịch", "Ngày", "Thành viên", "Loại", "Danh mục", "Nguồn tiền", "Số tiền (VNĐ)", "Ghi chú"])
    for item in report["items"]:
        writer.writerow([
            item.idExpense,
            item.date.strftime("%d/%m/%Y"),
            item.user.hoten,
            "Thu nhập" if item.category.type == "income" else "Chi tiêu",
            item.category.nameCategory,
            source_label(item.source),
            str(item.amount),
            item.purpose,
        ])
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename=bao_cao_chi_tieu_{report['start'].isoformat()}_{report['end'].isoformat()}.csv"})
