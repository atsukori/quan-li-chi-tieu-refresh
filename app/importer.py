from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import json
from pathlib import Path
import uuid

from openpyxl import load_workbook

from .extensions import db
from .models import Category, Expense, User
from .utils import VALID_SOURCES, find_duplicate_expense, parse_amount, parse_iso_date


REQUIRED_COLUMNS = {"idUser", "idCategory", "date", "amount", "purpose", "source"}


def _as_int(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        decimal_value = Decimal(str(value).strip())
        if not decimal_value.is_finite() or decimal_value != decimal_value.to_integral_value():
            return None
        number = int(decimal_value)
    except (InvalidOperation, TypeError, ValueError):
        return None
    return number if number > 0 else None


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_iso_date(str(value).strip())


def _stage_dir(instance_path):
    directory = Path(instance_path) / "imports"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def parse_expense_workbook(file_stream, instance_path, fallback_user_id):
    workbook = load_workbook(file_stream, read_only=True, data_only=True)
    try:
        sheet = next((item for item in workbook.worksheets if item.title.strip().lower() in {"giao dịch", "giao dich"}), None)
        if sheet is None:
            raise ValueError("Không tìm thấy sheet 'Giao dịch'.")

        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = {str(value).strip(): index for index, value in enumerate(header_values) if value is not None}
        missing = sorted(REQUIRED_COLUMNS - headers.keys())
        if missing:
            raise ValueError("Thiếu cột bắt buộc: " + ", ".join(missing))

        records = []
        seen = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            source = str(row[headers["source"]] or "").strip().lower()
            raw_user_id = row[headers["idUser"]]
            user_id = _as_int(raw_user_id)
            if not user_id and source in {"common", "unit"}:
                user_id = fallback_user_id
            category_id = _as_int(row[headers["idCategory"]])
            transaction_date = _as_date(row[headers["date"]])
            amount = parse_amount(row[headers["amount"]])
            purpose = str(row[headers["purpose"]] or "").strip()
            errors = []

            user = db.session.get(User, user_id) if user_id else None
            category = db.session.get(Category, category_id) if category_id else None
            if not user or not user.is_active_account:
                errors.append("Mã thành viên không tồn tại hoặc đang bị khóa.")
            if not category or not category.is_active:
                errors.append("Mã danh mục không tồn tại hoặc đang bị ẩn.")
            if not transaction_date:
                errors.append("Ngày không hợp lệ (dùng ngày Excel hoặc YYYY-MM-DD).")
            if amount is None:
                errors.append("Số tiền phải lớn hơn 0.")
            if source not in VALID_SOURCES:
                errors.append("Nguồn phải là personal, common hoặc unit.")
            if len(purpose) > 500:
                errors.append("Ghi chú vượt quá 500 ký tự.")

            record = {
                "row_number": row_number,
                "user_id": user_id,
                "user_name": user.hoten if user else "Không hợp lệ",
                "category_id": category_id,
                "category_name": category.nameCategory if category else "Không hợp lệ",
                "date": transaction_date.isoformat() if transaction_date else "",
                "amount": str(amount) if amount is not None else "",
                "purpose": purpose,
                "source": source,
                "status": "error" if errors else "valid",
                "message": "; ".join(errors),
            }
            if not errors:
                key = ((user_id if source == "personal" else None), transaction_date.isoformat(), str(amount), category_id, source)
                duplicate = find_duplicate_expense(user_id, transaction_date, amount, category_id, source)
                if duplicate or key in seen:
                    record["status"] = "duplicate"
                    record["message"] = "Giao dịch trùng trong database hoặc trong file; sẽ bỏ qua."
                seen.add(key)
            records.append(record)
            if len(records) > 5000:
                raise ValueError("Tệp vượt quá giới hạn 5.000 dòng giao dịch.")

        if not records:
            raise ValueError("File không có dòng giao dịch nào để nhập.")
        token = uuid.uuid4().hex
        stage_path = _stage_dir(instance_path) / f"{token}.json"
        stage_path.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        return token, records
    finally:
        workbook.close()


def read_staged_import(instance_path, token):
    if not token or len(token) != 32 or any(char not in "0123456789abcdef" for char in token):
        return None, None
    directory = _stage_dir(instance_path).resolve()
    path = (directory / f"{token}.json").resolve()
    if path.parent != directory or not path.is_file():
        return None, None
    try:
        if datetime.now().timestamp() - path.stat().st_mtime > 24 * 60 * 60:
            path.unlink(missing_ok=True)
            return None, None
        return path, json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None, None


def apply_staged_import(records):
    imported = 0
    skipped_duplicates = 0
    skipped_invalid = 0
    for record in records:
        if record.get("status") != "valid":
            if record.get("status") == "duplicate":
                skipped_duplicates += 1
            else:
                skipped_invalid += 1
            continue
        transaction_date = date.fromisoformat(record["date"])
        amount = parse_amount(record["amount"])
        if find_duplicate_expense(record["user_id"], transaction_date, amount, record["category_id"], record["source"]):
            skipped_duplicates += 1
            continue
        db.session.add(Expense(
            idUser=record["user_id"], idCategory=record["category_id"], date=transaction_date,
            amount=amount, purpose=record["purpose"], source=record["source"],
        ))
        imported += 1
    return imported, skipped_duplicates, skipped_invalid
